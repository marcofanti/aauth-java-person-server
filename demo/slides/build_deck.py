"""Rebuild aauth-java-person-server.pptx from slides.html.

Each slide is rendered to a 1920x1080 PNG with headless Chrome, then packed one image per
slide into a 16:9 deck. Run from the repo root:

    uv run --no-project --with python-pptx==1.0.2 demo/slides/build_deck.py
"""

import re
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
SLIDES_DIR = Path(__file__).resolve().parent
SOURCE = SLIDES_DIR / "slides.html"
OUTPUT = SLIDES_DIR / "aauth-java-person-server.pptx"
SLIDE_START = re.compile(r'<div class="slide(?: [^"]*)?"')
SHOW_ONE = (
    "body{margin:0;background:#f6f7f9} "
    ".slide{display:none;margin:0!important} .slide[data-show]{display:block}\n"
)


def render(work: Path) -> list[Path]:
    html = SOURCE.read_text(encoding="utf-8")
    images = []
    for number, match in enumerate(SLIDE_START.finditer(html), start=1):
        start = match.start()
        page = html[:start] + html[start:].replace(
            '<div class="slide', '<div data-show="1" class="slide', 1
        )
        page = page.replace("</style>", SHOW_ONE + "</style>", 1)
        page_file = work / f"slide{number:02d}.html"
        page_file.write_text(page, encoding="utf-8")
        image = work / f"slide{number:02d}.png"
        subprocess.run(
            [
                CHROME,
                "--headless=new",
                "--disable-gpu",
                "--hide-scrollbars",
                "--force-device-scale-factor=1",
                "--window-size=1920,1080",
                # Gives the Typekit web fonts time to load before the screenshot.
                "--virtual-time-budget=8000",
                f"--screenshot={image}",
                page_file.as_uri(),
            ],
            check=True,
            capture_output=True,
        )
        images.append(image)
    return images


def main() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        images = render(Path(tmp))
        deck = Presentation()
        deck.slide_width, deck.slide_height = Emu(12192000), Emu(6858000)
        for image in images:
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_picture(
                str(image), 0, 0, width=deck.slide_width, height=deck.slide_height
            )
        deck.save(OUTPUT)
    print(f"{len(images)} slides -> {OUTPUT}", file=sys.stderr)


if __name__ == "__main__":
    main()
